"""Builds the spreadsheet examples (3 levels per class) as .xlsx files that open in Google Planilhas.
Formulas are written with English names and commas (the .xlsx standard); Google Planilhas in PT-BR
shows them translated (SE, SOMA, PROCV...). The example text of each level comes from _data/infobasica.yml.
Uso: python3 _tools/infobasica/gerar_planilhas.py <raiz do site>"""
import sys, os, re, glob, yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import LineChart, ScatterChart, BarChart, Reference, Series
from openpyxl.chart.trendline import Trendline
from openpyxl.formatting.rule import CellIsRule

NIVEIS = {1: 'Básico', 2: 'Intermediário', 3: 'Avançado'}
ENTRADA = PatternFill('solid', fgColor='FFF2CC')   # células de entrada em amarelo claro
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
CHUVA = [320, 290, 360, 410, 380, 250, 190, 120, 80, 95, 140, 230]
TEMPS = [28, 31, 33, 25, 29, 35, 22]
DIAS = ['seg', 'ter', 'qua', 'qui', 'sex', 'sáb', 'dom']
GASTOS = [10, 12, 8, 15, 9, 20, 11]
MEDIDAS = [50.1, 49.7, 50.3, 49.9, 50.0, 50.25, 49.85, 50.15, 49.75, 50.05, 50.4, 49.95, 50.2, 49.6, 50.1]
ESTADOS = [('AC', 'Acre'), ('AL', 'Alagoas'), ('AP', 'Amapá'), ('AM', 'Amazonas'), ('BA', 'Bahia'), ('CE', 'Ceará'),
           ('DF', 'Distrito Federal'), ('ES', 'Espírito Santo'), ('GO', 'Goiás'), ('MA', 'Maranhão'), ('MT', 'Mato Grosso'),
           ('MS', 'Mato Grosso do Sul'), ('MG', 'Minas Gerais'), ('PA', 'Pará'), ('PB', 'Paraíba'), ('PR', 'Paraná'),
           ('PE', 'Pernambuco'), ('PI', 'Piauí'), ('RJ', 'Rio de Janeiro'), ('RN', 'Rio Grande do Norte'),
           ('RS', 'Rio Grande do Sul'), ('RO', 'Rondônia'), ('RR', 'Roraima'), ('SC', 'Santa Catarina'),
           ('SP', 'São Paulo'), ('SE', 'Sergipe'), ('TO', 'Tocantins')]


def coluna(ws, col, inicio, valores):
    for i, v in enumerate(valores):
        ws[f'{col}{inicio + i}'] = v

def cabecalho(ws, **celulas):
    for ref, texto in celulas.items():
        ws[ref] = texto
        ws[ref].font = Font(bold=True)

def entrada(ws, *refs):
    for r in refs:
        ws[r].fill = ENTRADA

def arraste(ws, col, inicio, fim, modelo):
    """Fills col[inicio..fim] with modelo, where {r} is the row and {a} the row above."""
    for r in range(inicio, fim + 1):
        ws[f'{col}{r}'] = modelo.format(r=r, a=r - 1)

def grafico(ws, tipo, titulo, x, series, ancora, eixo_x='', eixo_y='', tendencia=False, com_cabecalho=True):
    g = {'linha': LineChart, 'dispersao': ScatterChart, 'colunas': BarChart}[tipo]()
    g.title, g.x_axis.title, g.y_axis.title = titulo, eixo_x, eixo_y
    g.height, g.width = 8, 15
    rx = Reference(ws, range_string=f"'{ws.title}'!{x}")
    for faixa in series:                          # cada série com o cabeçalho na primeira linha
        ry = Reference(ws, range_string=f"'{ws.title}'!{faixa}")
        if tipo == 'dispersao':
            s = Series(ry, rx, title_from_data=com_cabecalho)
            s.marker.symbol = 'none'
            g.series.append(s)
        else:
            g.add_data(ry, titles_from_data=com_cabecalho)
    if tipo != 'dispersao':
        g.set_categories(rx)
    if tendencia:
        g.series[0].trendline = Trendline(trendlineType='linear')
    ws.add_chart(g, ancora)


# ---------- um construtor por (aula, nível) ----------
def a5(ws, k):
    cabecalho(ws, A1='Largura (m)', B1='Comprimento (m)', C1='Área (m²)')
    ws['A2'], ws['B2'], ws['C2'] = 4.5, 3.2, '=A2*B2'
    entrada(ws, 'A2', 'B2')
    if k >= 2:
        cabecalho(ws, D1='Perímetro (m)'); ws['D2'] = '=2*(A2+B2)'
    if k == 3:
        cabecalho(ws, E1='Preço do m² (R$)', F1='Custo com 10% de perda (R$)')
        ws['E2'], ws['F2'] = 59.9, '=C2*1.1*E2'
        ws['F2'].number_format = '"R$" #,##0.00'; entrada(ws, 'E2')

def a6(ws, k):
    if k == 1:
        cabecalho(ws, A1='Preço (R$)', A2='Duas unidades (R$)'); ws['B1'], ws['B2'] = 25, '=B1*2'; entrada(ws, 'B1')
    elif k == 2:
        cabecalho(ws, A1='Preço (R$)', A2='Desconto (%)', A3='Preço final (R$)')
        ws['B1'], ws['B2'], ws['B3'] = 100, 15, '=B1-B1*B2/100'; entrada(ws, 'B1', 'B2')
    else:
        cabecalho(ws, A1='Leitura anterior (kWh)', A2='Leitura atual (kWh)', A3='Preço do kWh (R$)', A4='Taxa fixa (R$)',
                  A5='Consumo (kWh)', A6='Valor da conta (R$)')
        ws['B1'], ws['B2'], ws['B3'], ws['B4'] = 1200, 1350, 0.85, 12
        ws['B5'], ws['B6'] = '=B2-B1', '=ROUND((B2-B1)*B3+B4,2)'; entrada(ws, 'B1', 'B2', 'B3', 'B4')

def a7(ws, k):
    if k == 1:
        cabecalho(ws, A1='Conta', B1='Resultado')
        for i, (t, f) in enumerate([('17+5', '=17+5'), ('17/5', '=17/5'), ('17^2', '=17^2'), ('MOD(17; 5)', '=MOD(17,5)')], 2):
            ws[f'A{i}'], ws[f'B{i}'] = t, f
    elif k == 2:
        cabecalho(ws, A1='Raio (m)', B1='Área (m²)', C1='Comprimento (m)')
        ws['A2'], ws['B2'], ws['C2'] = 2, '=ROUND(PI()*A2^2,2)', '=ROUND(2*PI()*A2,2)'; entrada(ws, 'A2')
    else:
        cabecalho(ws, A1='Segundos', B1='Horas', C1='Minutos', D1='Segundos restantes')
        ws['A2'], ws['B2'], ws['C2'], ws['D2'] = 3725, '=QUOTIENT(A2,3600)', '=QUOTIENT(MOD(A2,3600),60)', '=MOD(A2,60)'
        entrada(ws, 'A2')

def a8(ws, k):
    if k == 1:
        cabecalho(ws, A1='Idade', B1='Idade >= 16'); coluna(ws, 'A', 2, [15, 16, 17, 40]); arraste(ws, 'B', 2, 5, '=A{r}>=16')
    elif k == 2:
        cabecalho(ws, A1='Idade', B1='Voto facultativo (16 ou 17)', C1='Menor de 16 ou 70+')
        coluna(ws, 'A', 2, [14, 16, 17, 40, 72])
        arraste(ws, 'B', 2, 6, '=AND(A{r}>=16,A{r}<18)'); arraste(ws, 'C', 2, 6, '=OR(A{r}<16,A{r}>=70)')
    else:
        cabecalho(ws, A1='Ano', B1='Bissexto?'); coluna(ws, 'A', 2, [1900, 2000, 2024, 2025])
        arraste(ws, 'B', 2, 5, '=OR(AND(MOD(A{r},4)=0,MOD(A{r},100)<>0),MOD(A{r},400)=0)')

def a9(ws, k):
    if k == 1:
        cabecalho(ws, A1='Número', B1='Par ou ímpar'); coluna(ws, 'A', 2, [3, 8, 15, 22, 7, 10])
        arraste(ws, 'B', 2, 7, '=IF(MOD(A{r},2)=0,"Par","Ímpar")')
    elif k == 2:
        cabecalho(ws, A1='Estudante', B1='Nota 1', C1='Nota 2', D1='Média', E1='Situação')
        for i, (n, a, b) in enumerate([('Ana', 8, 7), ('Bia', 5, 6.5), ('Caio', 9, 9.5), ('Davi', 6, 8)], 2):
            ws[f'A{i}'], ws[f'B{i}'], ws[f'C{i}'] = n, a, b
        arraste(ws, 'D', 2, 5, '=(B{r}+C{r})/2'); arraste(ws, 'E', 2, 5, '=IF(D{r}>=7,"Aprovado","Reprovado")')
    else:
        cabecalho(ws, A1='Medida (mm)', B1='Situação'); coluna(ws, 'A', 2, MEDIDAS[:10])
        arraste(ws, 'B', 2, 11, '=IF(ABS(A{r}-50)<=0.2,"Aprovada","Reprovada")')
        ws.conditional_formatting.add('B2:B11', CellIsRule(operator='equal', formula=['"Reprovada"'],
                                      fill=PatternFill('solid', fgColor='F4CCCC'), font=Font(color='990000')))

def a10(ws, k):
    if k == 1:
        cabecalho(ws, A1='pH', B1='Classificação'); coluna(ws, 'A', 2, [2, 7, 9.5, 6.8, 12])
        arraste(ws, 'B', 2, 6, '=IF(A{r}<7,"Ácido",IF(A{r}=7,"Neutro","Básico"))')
    elif k == 2:
        cabecalho(ws, A1='Estudante', B1='Nota', C1='Conceito')
        coluna(ws, 'A', 2, ['Ana', 'Bia', 'Caio', 'Davi', 'Eva']); coluna(ws, 'B', 2, [9.5, 7.2, 5, 3.8, 8.9])
        arraste(ws, 'C', 2, 6, '=IFS(B{r}>=9,"A",B{r}>=7,"B",B{r}>=5,"C",TRUE,"D")')
    else:
        cabecalho(ws, A1='Velocidade (km/h)', B1='Limite (km/h)', C1='Excesso', D1='Multa')
        coluna(ws, 'A', 2, [55, 70, 85, 100]); coluna(ws, 'B', 2, [60] * 4)
        arraste(ws, 'C', 2, 5, '=(A{r}-B{r})/B{r}')
        arraste(ws, 'D', 2, 5, '=IFS(C{r}<=0,"Sem multa",C{r}<=20%,"Leve",C{r}<=50%,"Grave",TRUE,"Gravíssima")')
        for r in range(2, 6): ws[f'C{r}'].number_format = '0%'

def a13(ws, k):
    if k == 1:
        ws['A1'], ws['A2'] = 1, 2; entrada(ws, 'A1', 'A2')
        ws['C1'] = 'Selecione A1:A2 e arraste o quadradinho do canto até A10.'
    elif k == 2:
        cabecalho(ws, A1='Múltiplos de 5'); ws['A2'] = '=SEQUENCE(21,1,0,5)'
    else:
        cabecalho(ws, A1='Tabuada do', A2='Multiplicador', B2='Resultado'); ws['B1'] = 7; entrada(ws, 'B1')
        ws['A3'] = '=SEQUENCE(10)'; arraste(ws, 'B', 3, 12, '=$B$1*A{r}')

def a14(ws, k):
    if k <= 2:
        cabecalho(ws, A1='Dia', B1='Gasto (R$)'); coluna(ws, 'A', 2, DIAS); coluna(ws, 'B', 2, GASTOS)
        cabecalho(ws, A9='Total'); ws['B9'] = '=SUM(B2:B8)'
        if k == 2:
            cabecalho(ws, A10='Média', C1='Acumulado (R$)'); ws['B10'] = '=AVERAGE(B2:B8)'
            arraste(ws, 'C', 2, 8, '=SUM($B$2:B{r})')
    else:
        cabecalho(ws, A1='N', B1='N! (fatorial)', C1='Soma de 1 a N', D1='Pela fórmula N(N+1)/2')
        ws['A2'], ws['B2'], ws['C2'], ws['D2'] = 6, '=FACT(A2)', '=SUM(SEQUENCE(A2))', '=A2*(A2+1)/2'; entrada(ws, 'A2')

def a15(ws, k):
    cabecalho(ws, A1='Dia', B1='Temperatura (°C)', D1='Maior', D2='Menor')
    coluna(ws, 'A', 2, DIAS); coluna(ws, 'B', 2, TEMPS)
    ws['E1'], ws['E2'] = '=MAX(B2:B8)', '=MIN(B2:B8)'
    if k >= 2:
        cabecalho(ws, D3='Dias acima de 30 °C'); ws['E3'] = '=COUNTIF(B2:B8,">30")'
    if k == 3:
        ws.conditional_formatting.add('B2:B8', CellIsRule(operator='greaterThan', formula=['30'], fill=PatternFill('solid', fgColor='FCE5CD')))
        grafico(ws, 'linha', 'Temperaturas da semana', 'A2:A8', ['B1:B8'], 'G2', 'Dia', '°C')

def a16(ws, k):
    if k == 1:
        cabecalho(ws, A1='Mês', B1='Saldo (R$)'); coluna(ws, 'A', 2, list(range(0, 25)))
        ws['B2'] = 1000; entrada(ws, 'B2'); arraste(ws, 'B', 3, 26, '=B{a}*1.01')
    elif k == 2:
        cabecalho(ws, A1='Mês', B1='Saldo (R$)', C1='Dobrou?', D1='Taxa ao mês'); coluna(ws, 'A', 2, list(range(0, 81)))
        ws['B2'], ws['E1'] = 1000, 0.01; ws['E1'].number_format = '0.0%'; entrada(ws, 'B2', 'E1')
        arraste(ws, 'B', 3, 82, '=B{a}*(1+$E$1)'); arraste(ws, 'C', 3, 82, '=IF(B{r}>=2*$B$2,"Dobrou","")')
    else:
        cabecalho(ws, A1='Mês', B1='0,5% ao mês', C1='1% ao mês', D1='2% ao mês'); coluna(ws, 'A', 2, list(range(0, 81)))
        for c, taxa in (('B', 0.005), ('C', 0.01), ('D', 0.02)):
            ws[f'{c}2'] = 1000; arraste(ws, c, 3, 82, f'={c}{{a}}*(1+{taxa})')
        grafico(ws, 'linha', 'Saldo com três taxas', 'A2:A82', ['B1:B82', 'C1:C82', 'D1:D82'], 'F2', 'Mês', 'R$')

def a17(ws, k):
    if k == 1:
        cabecalho(ws, A1='Dado'); ws['A2'] = '=RANDBETWEEN(1,6)'
        ws['C1'] = 'Edite qualquer célula para sortear de novo.'
    elif k == 2:
        arraste(ws, 'A', 1, 1000, '=RANDBETWEEN(0,1)')
        cabecalho(ws, C1='% de caras (1)'); ws['D1'] = '=AVERAGE(A1:A1000)'; ws['D1'].number_format = '0.0%'
    else:
        arraste(ws, 'A', 1, 600, '=RANDBETWEEN(1,6)')
        coluna(ws, 'C', 1, list(range(1, 7))); arraste(ws, 'D', 1, 6, '=COUNTIF($A$1:$A$600,C{r})')
        grafico(ws, 'colunas', 'Frequência de cada face', 'C1:C6', ['D1:D6'], 'F2', 'Face', 'Vezes', com_cabecalho=False)

def a18(ws, k):
    cabecalho(ws, A1='Partida', B1='Tentativas'); coluna(ws, 'A', 2, list(range(1, 11)))
    coluna(ws, 'B', 2, [7, 5, 6, 4, 7, 3, 5, 6, 4, 5]); entrada(ws, *[f'B{r}' for r in range(2, 12)])
    if k >= 2:
        cabecalho(ws, D1='Média', D2='Mínimo', D3='Máximo')
        ws['E1'], ws['E2'], ws['E3'] = '=AVERAGE(B2:B11)', '=MIN(B2:B11)', '=MAX(B2:B11)'
    if k == 3:
        grafico(ws, 'linha', 'Tentativas por partida', 'A2:A11', ['B1:B11'], 'G2', 'Partida', 'Tentativas', tendencia=True)

def a19(ws, k):
    if k == 1:
        cabecalho(ws, A1='Número', B1='Múltiplo de 3 ou de 7'); coluna(ws, 'A', 2, list(range(1, 101)))
        arraste(ws, 'B', 2, 101, '=IF(OR(MOD(A{r},3)=0,MOD(A{r},7)=0),A{r},"")')
    else:
        ws['A1'], ws['A2'] = 0, 1
        for r in range(3, 21): ws[f'A{r}'] = f'=A{r-2}+A{r-1}'
        if k == 3:
            arraste(ws, 'B', 3, 20, '=A{r}/A{a}')
            cabecalho(ws, D1='MDC(48; 36)'); ws['E1'] = '=GCD(48,36)'
            grafico(ws, 'linha', 'Razão entre termos vizinhos (→ 1,618)', 'B3:B20', ['B3:B20'], 'G2', 'Termo', 'Razão', com_cabecalho=False)

def a20(ws, k):
    if k == 1:
        cabecalho(ws, A1='Metros', B1='Centímetros'); coluna(ws, 'A', 2, list(range(1, 11))); arraste(ws, 'B', 2, 11, '=A{r}*100')
    elif k == 2:
        cabecalho(ws, A1='°C', B1='°F'); ws['A2'] = '=SEQUENCE(11,1,0,10)'; arraste(ws, 'B', 2, 12, '=A{r}*9/5+32')
    else:
        cabecalho(ws, A1='Operação', B1='Valor (R$)', C1='Saldo (R$)')
        ops = [('Saldo inicial', 0), ('Depósito', 100), ('Saque', 30), ('Depósito', 50), ('Saque', 20)]
        for i, (o, v) in enumerate(ops, 2):
            ws[f'A{i}'], ws[f'B{i}'] = o, v
        ws['C2'] = 0; arraste(ws, 'C', 3, 6, '=C{a}+IF(A{r}="Depósito",B{r},-B{r})')

def a21(ws, k):
    if k == 1:
        cabecalho(ws, A1='x', B1='y = 2x + 1'); coluna(ws, 'A', 2, list(range(-10, 11))); arraste(ws, 'B', 2, 22, '=2*A{r}+1')
        grafico(ws, 'linha', 'y = 2x + 1', 'A2:A22', ['B1:B22'], 'D2', 'x', 'y')
    else:
        cabecalho(ws, A1='Tempo (s)', B1='Terra (m)'); coluna(ws, 'A', 2, [round(i * 0.1, 1) for i in range(31)])
        arraste(ws, 'B', 2, 32, '=9.81*A{r}^2/2'); series = ['B1:B32']
        if k == 3:
            cabecalho(ws, C1='Lua (m)'); arraste(ws, 'C', 2, 32, '=1.62*A{r}^2/2'); series.append('C1:C32')
        grafico(ws, 'dispersao', 'Queda livre', 'A2:A32', series, 'E2', 'Tempo (s)', 'Distância (m)')

def a22(ws, k):
    if k == 1:
        cabecalho(ws, A1='Hora', B1='Bactérias'); coluna(ws, 'A', 2, list(range(0, 11)))
        ws['B2'] = 100; entrada(ws, 'B2'); arraste(ws, 'B', 3, 12, '=B{a}*2')
    else:
        cabecalho(ws, A1='Minuto/hora', B1='Peça (°C)'); coluna(ws, 'A', 2, list(range(0, 60)))
        ws['B2'] = 200; entrada(ws, 'B2'); arraste(ws, 'B', 3, 61, '=B{a}-0.1*(B{a}-25)')
        if k == 2:
            cabecalho(ws, D1='Minutos com 40 °C ou mais'); ws['E1'] = '=COUNTIF(B2:B60,">=40")'
        else:
            cabecalho(ws, C1='Solução (%)'); ws['C2'] = 100; arraste(ws, 'C', 3, 61, '=C{a}*(1-8%)')
            grafico(ws, 'linha', 'Resfriamento e concentração', 'A2:A61', ['B1:B61', 'C1:C61'], 'E2', 'Passo', 'Valor')

def a23(ws, k):
    if k == 1:
        cabecalho(ws, A1='Valor', C1='Soma', C2='Média'); coluna(ws, 'A', 2, GASTOS)
        ws['D1'], ws['D2'] = '=SUM(A2:A8)', '=AVERAGE(A2:A8)'
    elif k == 2:
        a15(ws, 2)
    else:
        a22(ws, 1); grafico(ws, 'linha', 'Bactérias', 'A2:A12', ['B1:B12'], 'D2', 'Hora', 'Bactérias')

def a24(ws, k):
    if k == 1:
        cabecalho(ws, A1='Dia', B1='Gasto', C1='Acumulado'); ws['A2'] = '=SEQUENCE(7)'; coluna(ws, 'B', 2, GASTOS)
        arraste(ws, 'C', 2, 8, '=SUM($B$2:B{r})')
    elif k == 2:
        a17(ws, 3); ws._charts = []   # só a revisão das fórmulas, sem gráfico
    else:
        a21(ws, 3)

def a25(ws, k):
    if k == 1:
        cabecalho(ws, A1='Texto', B1='Maiúsculas', C1='Minúsculas', D1='Caracteres'); ws['A2'] = 'Concreto Armado'; entrada(ws, 'A2')
        ws['B2'], ws['C2'], ws['D2'] = '=UPPER(A2)', '=LOWER(A2)', '=LEN(A2)'
    else:
        cabecalho(ws, A1='Nome', B1='Sobrenome'); coluna(ws, 'A', 2, ['Ana', 'Caio', 'Bia']); coluna(ws, 'B', 2, ['Souza', 'Lima', 'Rocha'])
        if k == 2:
            cabecalho(ws, C1='3 primeiras letras', D1='Nome completo')
            arraste(ws, 'C', 2, 4, '=LEFT(A{r},3)'); arraste(ws, 'D', 2, 4, '=CONCATENATE(A{r}," ",B{r})')
        else:
            cabecalho(ws, C1='E-mail institucional'); arraste(ws, 'C', 2, 4, '=LOWER(A{r}&"."&B{r}&"@escola.br")')

def a26(ws, k):
    if k <= 2:
        cabecalho(ws, A1='Nomes', C1='Primeiro', C2='Quantidade'); coluna(ws, 'A', 2, ['Ana', 'Bia', 'Caio', 'Davi', 'Eva'])
        ws['D1'], ws['D2'] = '=A2', '=COUNTA(A2:A6)'
        if k == 2:
            cabecalho(ws, C3='3º elemento'); ws['D3'] = '=INDEX(A2:A6,3)'
    else:
        cabecalho(ws, A1='Mês', B1='Número do mês →'); coluna(ws, 'A', 2, ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho',
                                                                          'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro'])
        ws['C1'], ws['D1'] = 3, '=INDEX(A2:A13,C1)'; entrada(ws, 'C1')

def a27(ws, k):
    cabecalho(ws, A1='Mês', B1='Chuva (mm)', D1='Total', D2='Média', D3='Máximo'); coluna(ws, 'A', 2, MESES); coluna(ws, 'B', 2, CHUVA)
    ws['E1'], ws['E2'], ws['E3'] = '=SUM(B2:B13)', '=AVERAGE(B2:B13)', '=MAX(B2:B13)'
    if k >= 2:
        cabecalho(ws, D4='Mês mais chuvoso'); ws['E4'] = '=INDEX(A2:A13,MATCH(MAX(B2:B13),B2:B13,0))'
    if k == 3:
        cabecalho(ws, D5='Meses acima da média', D6='1º maior', D7='2º maior', D8='3º maior')
        ws['E5'] = '=COUNTIF(B2:B13,">"&AVERAGE(B2:B13))'
        for i in (1, 2, 3): ws[f'E{5 + i}'] = f'=LARGE(B2:B13,{i})'

def a28(ws, k):
    if k <= 2:
        cabecalho(ws, A1='Livros'); coluna(ws, 'A', 2, ['Dom Casmurro', 'Vidas Secas', 'O Cortiço', 'Iracema', 'Capitães da Areia'])
        cabecalho(ws, C1='Tem Iracema?'); ws['D1'] = '=COUNTIF(A2:A20,"Iracema")>0'
        if k == 2:
            cabecalho(ws, C2='Posição de Iracema'); ws['D2'] = '=MATCH("Iracema",A2:A20,0)'
    else:
        cabecalho(ws, A1='Sigla', B1='Estado', C1='Sigla buscada →')
        for i, (s, n) in enumerate(ESTADOS, 2):
            ws[f'A{i}'], ws[f'B{i}'] = s, n
        ws['D1'], ws['E1'] = 'BA', '=VLOOKUP(D1,A2:B28,2,FALSE)'; entrada(ws, 'D1')

def a29(ws, k):
    if k == 1:
        cabecalho(ws, A1='Números', C1='Só os pares')
        coluna(ws, 'A', 2, [3, 8, 12, 5, 20, 7, 14, 9, 30, 11, 6, 17, 22, 1, 40, 13, 18, 25, 4])
        ws['C2'] = '=FILTER(A2:A20,MOD(A2:A20,2)=0)'
    else:
        cabecalho(ws, A1='Medida (mm)'); coluna(ws, 'A', 2, MEDIDAS)
        if k == 2:
            cabecalho(ws, C1='Fora da tolerância'); ws['C2'] = '=FILTER(A2:A16,ABS(A2:A16-50)>0.2)'
        else:
            cabecalho(ws, B1='Em polegadas', D1='% reprovado'); arraste(ws, 'B', 2, 16, '=ROUND(A{r}/25.4,3)')
            ws['E1'] = '=(COUNTIF(A2:A16,">50.2")+COUNTIF(A2:A16,"<49.8"))/COUNT(A2:A16)'; ws['E1'].number_format = '0%'

def a30(ws, k):
    if k == 1:
        cabecalho(ws, A1='Valor', B1='Arredondado', C1='Ao quadrado'); ws['A2'] = 3.14159; entrada(ws, 'A2')
        ws['B2'], ws['C2'] = '=ROUND(A2,2)', '=POWER(A2,2)'
    else:
        cabecalho(ws, A1='Largura (m)', B1='Comprimento (m)', C1='Resultado esperado (m²)', D1='Sua função nomeada →')
        ws['A2'], ws['B2'] = 5, 4; entrada(ws, 'A2', 'B2')
        if k == 2:
            ws['C2'] = '=A2*B2'
            ws['A4'] = 'Crie AREA_RET(largura; comprimento) em Dados → Funções nomeadas e escreva =AREA_RET(A2; B2) em E2.'
        else:
            cabecalho(ws, A3='Raio da coluna (m)'); ws['B3'] = 0.3; entrada(ws, 'B3')
            ws['C2'] = '=A2*B2-PI()*B3^2'
            ws['A5'] = 'Crie também AREA_CIRC(raio) e escreva =AREA_RET(5; 4)-AREA_CIRC(0,3) em E2.'

def a32(ws, k):
    if k == 1:
        cabecalho(ws, A1='Material', B1='Preço (R$)', D1='cimento')
        for i, (m, p) in enumerate([('cimento', 38.9), ('areia', 120), ('brita', 135), ('cal', 22.5), ('tijolo', 1.2)], 2):
            ws[f'A{i}'], ws[f'B{i}'] = m, p
        ws['E1'] = '=VLOOKUP("cimento",A2:B6,2,FALSE)'
    else:
        cabecalho(ws, A1='Produto', B1='Valor vendido (R$)')
        vendas = [('cimento', 38.9), ('areia', 120), ('cimento', 77.8), ('brita', 135), ('areia', 60), ('cimento', 38.9), ('cal', 45)]
        for i, (m, v) in enumerate(vendas, 2):
            ws[f'A{i}'], ws[f'B{i}'] = m, v
        cabecalho(ws, D1='Total de cimento'); ws['E1'] = '=SUMIF(A2:A20,"cimento",B2:B20)'
        if k == 3:
            cabecalho(ws, D3='Produto', E3='Total (confira com a tabela dinâmica)')
            for i, m in enumerate(['areia', 'brita', 'cal', 'cimento'], 4):
                ws[f'D{i}'], ws[f'E{i}'] = m, f'=SUMIF($A$2:$A$20,D{i},$B$2:$B$20)'

def a34(ws, k):
    if k == 1:
        cabecalho(ws, A1='Nome', B1='Nota')
        for i, (n, v) in enumerate([('Ana', 7.5), ('Bia', 9), ('Caio', 5.5), ('Davi', 8)], 2):
            ws[f'A{i}'], ws[f'B{i}'] = n, v
    elif k == 2:
        cabecalho(ws, A1='Trimestre', B1='Consumo medio')
        for i, v in enumerate([13.0, 17.7, 18.7, 13.0], 2):
            ws[f'A{i}'], ws[f'B{i}'] = i - 1, v
    else:
        cabecalho(ws, A1='Mês', B1='Consumo (m³)', D1='Trimestre', E1='Média (m³)')
        coluna(ws, 'A', 2, list(range(1, 13))); coluna(ws, 'B', 2, [12, 14, 13, 15, 18, 20, 21, 19, 16, 14, 13, 12])
        for t in range(4):
            ws[f'D{t + 2}'], ws[f'E{t + 2}'] = t + 1, f'=ROUND(AVERAGE(B{2 + 3 * t}:B{4 + 3 * t}),1)'
        grafico(ws, 'colunas', 'Consumo médio por trimestre', 'D2:D5', ['E1:E5'], 'G2', 'Trimestre', 'm³')


CONSTRUTORES = {5: a5, 6: a6, 7: a7, 8: a8, 9: a9, 10: a10, 13: a13, 14: a14, 15: a15, 16: a16, 17: a17, 18: a18,
                19: a19, 20: a20, 21: a21, 22: a22, 23: a23, 24: a24, 25: a25, 26: a26, 27: a27, 28: a28,
                29: a29, 30: a30, 32: a32, 34: a34}

# Funções novas precisam do prefixo do Excel no .xlsx; o Google Planilhas aceita _xlfn.IFS/_xlfn.SEQUENCE,
# mas não _xlfn._xlws.FILTER, por isso FILTER fica sem prefixo
PREFIXOS = {'IFS(': '_xlfn.IFS(', 'SEQUENCE(': '_xlfn.SEQUENCE(', 'FILTER(': 'FILTER('}

def prefixar(ws):
    for linha in ws.iter_rows():
        for c in linha:
            if isinstance(c.value, str) and c.value.startswith('='):
                for nome, novo in PREFIXOS.items():
                    c.value = re.sub(r'(?<![\w.])' + re.escape(nome), novo, c.value)

def texto_simples(md):
    # tira a marcação do Markdown, mas mantém o * das fórmulas (que ficam entre crases)
    partes = md.split('`')
    return ''.join(p if i % 2 else p.replace('*', '') for i, p in enumerate(partes))

if __name__ == '__main__':
    raiz = sys.argv[1]
    aulas = {a['n']: a for u in yaml.safe_load(open(os.path.join(raiz, '_data/infobasica.yml'))) for a in u['aulas']}
    destino = os.path.join(raiz, 'assets/infobasica/planilhas')
    os.makedirs(destino, exist_ok=True)
    for antigo in glob.glob(os.path.join(destino, '*.xlsx')):
        os.remove(antigo)
    total = 0
    for n, construtor in CONSTRUTORES.items():
        assert len(aulas[n]['planilhas']) == 3, n
        for k in (1, 2, 3):
            wb = Workbook()
            ws = wb.active
            ws.title = 'Exemplo'
            construtor(ws, k)
            prefixar(ws)
            for col in 'ABCDEF':
                ws.column_dimensions[col].width = 18
            leia = wb.create_sheet('Leia-me')
            leia['A1'] = f"Aula {n} – {aulas[n]['tema']}"
            leia['A1'].font = Font(bold=True, size=13)
            leia['A2'] = f'Nível {k} ({NIVEIS[k]})'
            leia['A2'].font = Font(bold=True)
            leia['A4'] = texto_simples(aulas[n]['planilhas'][k - 1])
            leia['A4'].data_type = 's'   # texto, mesmo quando começa com =
            leia['A6'] = 'Células em amarelo são entradas: mude os valores e veja as fórmulas recalcularem.'
            leia['A7'] = 'Esta é a sua cópia: altere à vontade.'
            leia.column_dimensions['A'].width = 120
            wb.save(os.path.join(destino, f'aula{n:02d}_nivel{k}.xlsx'))
            total += 1
    print(total, 'planilhas')
